#!/usr/bin/env python3
"""Build Purpose Academy product overview PowerPoint (16:9, real screenshots)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Prefer local grant-build shots; fall back to Cursor artifacts copies.
SHOT_DIRS = [
    Path("/tmp/grant-build/shots"),
    Path("/opt/cursor/artifacts"),
]
OUT = Path("/workspace/docs/grant/purpose-academy-overview.pptx")
ARTIFACT_COPY = Path("/opt/cursor/artifacts/purpose-academy-overview.pptx")

# Product palette — matches src/styles.css (navy + amber, not purple AI cliché)
NAVY = RGBColor(0x0B, 0x2F, 0x5C)
DEEP = RGBColor(0x04, 0x15, 0x26)
INK = RGBColor(0x07, 0x1E, 0x3D)
MUTED = RGBColor(0x2F, 0x41, 0x56)
AMBER = RGBColor(0xE8, 0xA3, 0x17)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF3, 0xF6, 0xFB)
SOFT = RGBColor(0xD8, 0xE0, 0xEC)


def resolve(name: str) -> Path:
    """Resolve a shot by exact name, then common grant_* aliases."""
    aliases = {
        "01-mode-chooser.png": ["grant_mode_chooser.png"],
        "02-splash.png": ["grant_mode_chooser.png"],
        "03b-toolbox-trays.png": ["grant_toolbox.png", "03-toolbox.png"],
        "04-practice-hub.png": ["grant_practice_hub.png"],
        "05-passport.png": ["grant_skills_passport.png"],
        "06-visual-vocab.png": ["grant_visual_vocab.png"],
        "07-eye-spy.png": ["grant_eye_spy.png"],
        "09-dignity.png": ["10-dignity-chrome.png"],
        "phone-home.png": [
            "demo-phone-home-full.png",
            "live-phone-home-icon.png",
            "responsive-open-phone.png",
        ],
        "phone-vv.png": [
            "walkthrough-vv-phone.png",
            "sprint1-phone-vv.png",
            "phone-bcs-readable.png",
        ],
    }
    candidates = [name] + aliases.get(name, [])
    for directory in SHOT_DIRS:
        for candidate in candidates:
            path = directory / candidate
            if path.is_file() and path.stat().st_size > 10_000:
                return path
    raise SystemExit(f"Missing screenshot for {name} (tried {candidates})")


# Product overview — short captions, no fake stats, Calgary voice.
SLIDES = [
    {
        "kind": "title",
        "title": "Purpose Academy",
        "kicker": "Product overview · Calgary trades school learning app",
        "body": "Workplace English and trades foundations — Construction, Logistics, and Community Support — practised one station at a time.",
        "image": "01-mode-chooser.png",
        "caption": "Desktop or Phone — same school, two ways in.",
    },
    {
        "kind": "text_image",
        "title": "Who it’s for",
        "kicker": "Adult learners first",
        "body": "Newcomers and first-time trades learners who want workplace English with adult dignity. Soft chrome for people who may fear school — never childish gamification.",
        "image": "02-splash.png",
        "caption": "Night-stage door — open into the workshop.",
    },
    {
        "kind": "dual",
        "title": "Desktop and phone",
        "kicker": "One school · two surfaces",
        "body": "Desktop opens the toolbox, then the full web course on a calm night stage. Phone is a live device-frame demo of the same path.",
        "image_left": "03b-toolbox-trays.png",
        "image_right": "phone-home.png",
        "caption_left": "Toolbox trays",
        "caption_right": "Phone home",
    },
    {
        "kind": "text_image",
        "title": "Three pathways",
        "kicker": "Construction · Logistics · Community Support",
        "body": "Same Practice spine for every door. Specialty vocabulary and scenes change with the pathway. Today’s station stays in focus.",
        "image": "04-practice-hub.png",
        "caption": "Practice hub — pick a pathway, enter today’s station.",
    },
    {
        "kind": "text_image",
        "title": "Visual Vocabulary",
        "kicker": "See → hear home language → hear English → type",
        "body": "Look at the object. Press a language you trust. Hear the English workplace word. Then type it. Finite pages — leave when you need to.",
        "image": "06-visual-vocab.png",
        "caption": "VocabSheet on the desktop night stage.",
    },
    {
        "kind": "dual",
        "title": "Stations that stay clear",
        "kicker": "Practice on phone and desktop",
        "body": "The same Visual Vocabulary station fits inside the phone frame. Task, picture, and languages stay unmistakable — no cluttered exam board.",
        "image_left": "06-visual-vocab.png",
        "image_right": "phone-vv.png",
        "caption_left": "Desktop VocabSheet",
        "caption_right": "Phone station",
    },
    {
        "kind": "text_image",
        "title": "Eye Spy",
        "kicker": "Clear the board",
        "body": "Tap a picture, name it in English, clear the board. Same workplace words in busy shop scenes — a concrete win without scoreboard shame.",
        "image": "07-eye-spy.png",
        "caption": "Eye Spy — match the English name, clear every pair.",
    },
    {
        "kind": "text_image",
        "title": "Learner dignity",
        "kicker": "Finite path · Done for now",
        "body": "One station in focus. The Practice banner names the path. Done for now is always available. No casino streaks. No “prove yourself” exam theatre in learner chrome.",
        "image": "09-dignity.png",
        "caption": "Soft Practice chrome with Done for now.",
    },
    {
        "kind": "text_image",
        "title": "Skills Passport",
        "kicker": "Honest continuum — not a guarantee",
        "body": "Learned → Practised → Competent. Practice alone is not Red Seal, apprenticeship certification, or a job guarantee. Competent still needs a real instructor.",
        "image": "05-passport.png",
        "caption": "Skills Passport honesty copy in Practice.",
    },
    {
        "kind": "close",
        "title": "Built in Calgary",
        "kicker": "Purpose Academy Application",
        "body": "Construction · Logistics · Community Support Foundations. Open the Contact tray in the toolbox when you want to talk next.",
        "image": "01-mode-chooser.png",
        "caption": "One school. One station at a time.",
    },
]


def set_run(run, *, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size=18,
    bold=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    font="Calibri",
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def fill_slide_bg(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def fit_picture(slide, path: Path, left, top, max_w, max_h):
    """Place a picture scaled to fit inside a box without cropping."""
    from PIL import Image

    with Image.open(path) as im:
        w_px, h_px = im.size
    aspect = w_px / h_px
    box_aspect = float(max_w) / float(max_h)
    if aspect >= box_aspect:
        width = max_w
        height = max_w / aspect
    else:
        height = max_h
        width = max_h * aspect
    left_adj = left + (max_w - width) / 2
    top_adj = top + (max_h - height) / 2
    return slide.shapes.add_picture(str(path), left_adj, top_adj, width=width, height=height)


def add_footer(slide, index: int, total: int, caption: str, *, light=False):
    color = SOFT if light else MUTED
    add_textbox(
        slide,
        Inches(0.5),
        Inches(7.05),
        Inches(12.3),
        Inches(0.3),
        f"{index} / {total}  ·  {caption}",
        size=11,
        color=color,
    )


def build_title(slide, spec, index, total):
    fill_slide_bg(slide, DEEP)
    add_textbox(
        slide,
        Inches(0.6),
        Inches(0.4),
        Inches(12),
        Inches(0.35),
        spec["kicker"].upper(),
        size=13,
        bold=True,
        color=AMBER,
    )
    add_textbox(
        slide,
        Inches(0.6),
        Inches(0.8),
        Inches(12),
        Inches(0.7),
        spec["title"],
        size=42,
        bold=True,
        color=WHITE,
        font="Georgia",
    )
    add_textbox(
        slide,
        Inches(0.6),
        Inches(1.55),
        Inches(12),
        Inches(0.65),
        spec["body"],
        size=16,
        color=SOFT,
    )
    fit_picture(
        slide,
        resolve(spec["image"]),
        Inches(1.1),
        Inches(2.35),
        Inches(11.1),
        Inches(4.4),
    )
    add_footer(slide, index, total, spec["caption"], light=True)


def build_text_image(slide, spec, index, total):
    fill_slide_bg(slide, PAPER)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.35),
        Inches(5.9),
        Inches(0.35),
        spec["kicker"].upper(),
        size=12,
        bold=True,
        color=AMBER,
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.7),
        Inches(5.9),
        Inches(0.7),
        spec["title"],
        size=30,
        bold=True,
        color=NAVY,
        font="Georgia",
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.55),
        Inches(5.7),
        Inches(3.8),
        spec["body"],
        size=17,
        color=INK,
    )
    fit_picture(
        slide,
        resolve(spec["image"]),
        Inches(6.5),
        Inches(0.45),
        Inches(6.4),
        Inches(6.3),
    )
    add_footer(slide, index, total, spec["caption"])


def build_dual(slide, spec, index, total):
    fill_slide_bg(slide, PAPER)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.28),
        Inches(12.3),
        Inches(0.3),
        spec["kicker"].upper(),
        size=12,
        bold=True,
        color=AMBER,
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.55),
        Inches(12.3),
        Inches(0.55),
        spec["title"],
        size=28,
        bold=True,
        color=NAVY,
        font="Georgia",
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(1.15),
        Inches(12.3),
        Inches(0.55),
        spec["body"],
        size=15,
        color=INK,
    )
    fit_picture(
        slide,
        resolve(spec["image_left"]),
        Inches(0.45),
        Inches(1.85),
        Inches(6.1),
        Inches(4.7),
    )
    fit_picture(
        slide,
        resolve(spec["image_right"]),
        Inches(6.75),
        Inches(1.85),
        Inches(6.1),
        Inches(4.7),
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(6.65),
        Inches(6.0),
        Inches(0.3),
        spec["caption_left"],
        size=12,
        color=MUTED,
    )
    add_textbox(
        slide,
        Inches(6.8),
        Inches(6.65),
        Inches(6.0),
        Inches(0.3),
        spec["caption_right"],
        size=12,
        color=MUTED,
    )
    add_footer(
        slide,
        index,
        total,
        f"{spec['caption_left']}  ·  {spec['caption_right']}",
    )


def build_close(slide, spec, index, total):
    fill_slide_bg(slide, DEEP)
    add_textbox(
        slide,
        Inches(0.6),
        Inches(0.55),
        Inches(12),
        Inches(0.35),
        spec["kicker"].upper(),
        size=13,
        bold=True,
        color=AMBER,
    )
    add_textbox(
        slide,
        Inches(0.6),
        Inches(1.0),
        Inches(12),
        Inches(0.7),
        spec["title"],
        size=40,
        bold=True,
        color=WHITE,
        font="Georgia",
    )
    add_textbox(
        slide,
        Inches(0.6),
        Inches(1.85),
        Inches(12),
        Inches(0.9),
        spec["body"],
        size=18,
        color=SOFT,
    )
    fit_picture(
        slide,
        resolve(spec["image"]),
        Inches(2.0),
        Inches(2.95),
        Inches(9.3),
        Inches(3.7),
    )
    add_footer(slide, index, total, spec["caption"], light=True)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    total = len(SLIDES)

    builders = {
        "title": build_title,
        "text_image": build_text_image,
        "dual": build_dual,
        "close": build_close,
    }

    for i, spec in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        builders[spec["kind"]](slide, spec, i, total)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    ARTIFACT_COPY.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(ARTIFACT_COPY))
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")
    print(f"wrote {ARTIFACT_COPY}")
    print("slides:", " · ".join(s["title"] for s in SLIDES))


if __name__ == "__main__":
    main()
