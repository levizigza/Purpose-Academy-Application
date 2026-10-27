#!/usr/bin/env python3
"""Build Purpose Academy grant overview PowerPoint (desktop screenshots + short captions)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

SHOTS = Path("/tmp/grant-build/shots")
OUT = Path("/opt/cursor/artifacts/purpose-academy-grant-overview.pptx")
REPO_COPY = Path("/workspace/docs/grant/purpose-academy-grant-overview.pptx")

# Brand-ish palette (navy + amber — matches product, not purple AI cliché)
NAVY = RGBColor(0x0B, 0x2F, 0x5C)
INK = RGBColor(0x1A, 0x1F, 0x2E)
MUTED = RGBColor(0x4A, 0x55, 0x68)
AMBER = RGBColor(0xC9, 0x8A, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF7, 0xF5, 0xF0)

SLIDES = [
    {
        "title": "Purpose Academy",
        "kicker": "Grant overview · Calgary trades school learning app",
        "body": "A concise look at what exists today: three pathways, Practice spine, desktop night stage, and honest progress toward employment readiness.",
        "image": "01-mode-chooser.png",
        "caption": "Desktop or Phone — learner workshop entry.",
    },
    {
        "title": "Who it’s for",
        "kicker": "Newcomers & first-time trades learners",
        "body": "Adults who want workplace English and trades foundations without exam theatre. Soft learner-fear chrome. Adult dignity — no childish gamification.",
        "image": "02-splash.png",
        "caption": "Night-stage opening: three doors into work.",
    },
    {
        "title": "Three pathways",
        "kicker": "One school · three doors",
        "body": "Construction · Logistics · Community Support Foundations. Same Practice spine; specialty vocabulary and scenes per pathway.",
        "image": "04-practice-hub.png",
        "caption": "Practice hub — choose pathway, one station at a time.",
    },
    {
        "title": "Practice loop",
        "kicker": "See → hear home language → hear English → type",
        "body": "Visual Vocabulary: look at the object, press a language you trust, hear the English workplace word, then type. Finite stations. Done for now when you need to stop.",
        "image": "06-visual-vocab.png",
        "caption": "VocabSheet on desktop night stage + paper card.",
    },
    {
        "title": "Clear the board",
        "kicker": "English Eye Spy",
        "body": "Tap a picture, name it in English, clear the board. Same words, busy shop scenes — concrete win without scoreboard shame.",
        "image": "07-eye-spy.png",
        "caption": "Eye Spy exercise mode — clear every pair.",
    },
    {
        "title": "Desktop experience",
        "kicker": "Toolbox → Practice / Student / Instructor / Admin",
        "body": "Desktop mode opens the toolbox trays, then the full web course on a calm night stage. Phone mode is a live device-frame demo of the same school.",
        "image": "03b-toolbox-trays.png",
        "caption": "Tool Chest trays — Practice highlighted.",
    },
    {
        "title": "Learner dignity",
        "kicker": "Soft chrome · finite path",
        "body": "One station in focus. Practice banner names the finite path. Done for now is always available. We refuse casino streaks and “prove yourself” exam language in learner chrome.",
        "image": "09-dignity.png",
        "caption": "Practice surface with Done for now / soft chrome.",
    },
    {
        "title": "Progress & Skills Passport",
        "kicker": "Honest continuum — not a guarantee",
        "body": "Learned → Practised → Competent. Practice alone is not Red Seal, apprenticeship certification, or a job guarantee. Competent still needs a real instructor.",
        "image": "05-passport.png",
        "caption": "Skills Passport honesty copy in Practice.",
    },
    {
        "title": "Built in Calgary",
        "kicker": "Contact & next step",
        "body": "Purpose Academy Application — Construction, Logistics, Community Support Foundations. For partnership or grant follow-up, use the Contact tray in the toolbox or the school’s public contact channel.",
        "image": "08-programs.png",
        "caption": "Programs surface — three pathways, honest depth.",
    },
]


def set_run(run, *, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(slide, left, top, width, height, text, *, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def fill_slide_bg(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    for i, spec in enumerate(SLIDES):
        slide = prs.slides.add_slide(blank)
        fill_slide_bg(slide, PAPER if i > 0 else NAVY)

        img_path = SHOTS / spec["image"]
        if not img_path.exists():
            raise SystemExit(f"Missing image {img_path}")

        if i == 0:
            add_textbox(
                slide,
                Inches(0.6),
                Inches(0.45),
                Inches(12),
                Inches(0.4),
                spec["kicker"].upper(),
                size=14,
                bold=True,
                color=AMBER,
            )
            add_textbox(
                slide,
                Inches(0.6),
                Inches(0.85),
                Inches(12),
                Inches(0.7),
                spec["title"],
                size=40,
                bold=True,
                color=WHITE,
            )
            add_textbox(
                slide,
                Inches(0.6),
                Inches(1.55),
                Inches(12),
                Inches(0.7),
                spec["body"],
                size=16,
                color=RGBColor(0xD8, 0xE0, 0xEC),
            )
            # Screenshot lower
            slide.shapes.add_picture(str(img_path), Inches(1.2), Inches(2.4), width=Inches(10.9))
            add_textbox(
                slide,
                Inches(0.6),
                Inches(7.05),
                Inches(12),
                Inches(0.3),
                spec["caption"],
                size=12,
                color=RGBColor(0xB8, 0xC4, 0xD4),
            )
        else:
            add_textbox(
                slide,
                Inches(0.5),
                Inches(0.28),
                Inches(6.2),
                Inches(0.35),
                spec["kicker"].upper(),
                size=12,
                bold=True,
                color=AMBER,
            )
            add_textbox(
                slide,
                Inches(0.5),
                Inches(0.55),
                Inches(6.2),
                Inches(0.55),
                spec["title"],
                size=28,
                bold=True,
                color=NAVY,
            )
            add_textbox(
                slide,
                Inches(0.5),
                Inches(1.2),
                Inches(6.0),
                Inches(2.4),
                spec["body"],
                size=16,
                color=INK,
            )
            add_textbox(
                slide,
                Inches(0.5),
                Inches(6.9),
                Inches(6.0),
                Inches(0.35),
                f"{i + 1} / {len(SLIDES)}  ·  {spec['caption']}",
                size=11,
                color=MUTED,
            )
            # Right-side screenshot
            slide.shapes.add_picture(str(img_path), Inches(6.7), Inches(0.55), width=Inches(6.2))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    REPO_COPY.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(REPO_COPY))
    print(f"wrote {OUT}")
    print(f"wrote {REPO_COPY}")
    print("slides:", " · ".join(s["title"] for s in SLIDES))


if __name__ == "__main__":
    main()
